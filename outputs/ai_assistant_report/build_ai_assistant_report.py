from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "outputs/ai_assistant_report/AI_Assistant_Implementation_Report.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(90, 99, 112)
LIGHT_FILL = "F4F6F9"
GRID = "D9E2EF"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), GRID)

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def style_paragraph(p, before=0, after=8, line=1.333, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align


def add_para(doc, text="", style=None, bold=False, italic=False, color=None, size=11, after=8, before=0, align=None):
    p = doc.add_paragraph(style=style)
    style_paragraph(p, before=before, after=after, align=align)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color or INK, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0].cells
    for i, title in enumerate(headers):
        set_cell_shading(hdr[i], LIGHT_FILL)
        p = hdr[i].paragraphs[0]
        style_paragraph(p, after=0, line=1.1)
        run = p.add_run(title)
        set_run_font(run, size=10, color=DARK_BLUE, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            style_paragraph(p, after=0, line=1.1)
            run = p.add_run(str(value))
            set_run_font(run, size=9.5, color=INK)
    set_table_geometry(table, widths)
    add_para(doc, "", after=2)
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F9FC")
    p = cell.paragraphs[0]
    style_paragraph(p, after=0, line=1.1)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        run = p.add_run(line)
        set_run_font(run, name="Consolas", size=9.5, color=INK)
    add_para(doc, "", after=2)


def set_document_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1


def add_footer(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(footer, after=0, line=1.0)
    run = footer.add_run("Performance Evaluation System | AI Assistant Implementation Report")
    set_run_font(run, size=8.5, color=MUTED)


def add_cover(doc):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(header, after=0, line=1.0)
    run = header.add_run("Faculty of Engineering, University of Sri Jayewardenepura")
    set_run_font(run, size=9, color=MUTED)
    add_footer(section)

    add_para(doc, "Performance Evaluation System", bold=True, color=MUTED, size=12, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    title = doc.add_paragraph()
    style_paragraph(title, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = title.add_run("AI Assistant Implementation Report")
    set_run_font(run, size=24, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=22, line=1.15, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = subtitle.add_run("A practical design for a faculty-aware, data-grounded academic guidance chatbot")
    set_run_font(run, size=13.5, color=MUTED)

    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [4680, 4680])
    rows = [
        ("Prepared for", "CO3554 Data Management Project"),
        ("Prepared by", "K. M. L. N. Senadheera | 22/ENG/079 | EN108953"),
        ("Source basis", "Project proposal report, Faculty Handbook 2026, and current Supabase-centered PES design"),
        ("Date", "July 2026"),
    ]
    for r, (label, value) in enumerate(rows):
        for c in range(2):
            set_cell_shading(table.cell(r, c), LIGHT_FILL if r == 0 else "FFFFFF")
        p_label = table.cell(r, 0).paragraphs[0]
        p_value = table.cell(r, 1).paragraphs[0]
        style_paragraph(p_label, after=0, line=1.1)
        style_paragraph(p_value, after=0, line=1.1)
        set_run_font(p_label.add_run(label), size=10, color=DARK_BLUE, bold=True)
        set_run_font(p_value.add_run(value), size=10, color=INK)
    add_para(doc, "", after=14)

    add_para(
        doc,
        "This report converts the proposed AI assistant feature into an implementation-ready design. It explains the correct system architecture, data foundation, query workflows, GPA planning logic, security model, evaluation method, and phased roadmap needed to build a trustworthy assistant for engineering students and faculty users.",
        size=11,
        after=8,
    )
    doc.add_page_break()


def build_doc():
    doc = Document()
    set_document_styles(doc)
    add_cover(doc)

    add_heading(doc, "1. Executive Summary", 1)
    add_para(
        doc,
        "The AI assistant should be implemented as a controlled academic guidance layer on top of the existing Performance Evaluation System. Its purpose is not only to chat, but to answer faculty-specific questions using the official handbook, the structured Supabase database, and deterministic academic calculations. This means the assistant must know the student's department, batch, completed results, remaining curriculum, attendance status, and faculty rules before giving personalized advice.",
    )
    add_para(
        doc,
        "The recommended design is a hybrid architecture: a Large Language Model for conversation, a Retrieval-Augmented Generation knowledge base for handbook explanations, Supabase tools for live student data, and verified calculation services for GPA, CGPA, eligibility, and target planning. This approach gives the system the flexibility of an AI chatbot while keeping academic answers grounded in authoritative data.",
    )

    add_heading(doc, "2. Project Context", 1)
    add_para(
        doc,
        "The current project proposal identifies the main weakness in the existing academic experience: students do not have a centralized, intelligent way to understand their performance, attendance eligibility, course progress, future semester subjects, minor/elective decisions, or risk level. The AI assistant is the feature that can connect these separate parts into one natural interface.",
    )
    add_para(
        doc,
        "The Faculty Handbook 2026 is the official rule source for departments, course codes, semester structures, credits, GPA contribution, examination rules, Dean's List criteria, and degree classification. Supabase is the operational source for each student's actual records. The assistant must combine both sources carefully: the handbook explains the rules, while Supabase provides the student's personal facts.",
    )

    add_heading(doc, "3. Primary Goals", 1)
    goals = [
        "Answer student questions about courses, semesters, credits, departments, GPA rules, repeat/resit rules, Dean's List eligibility, and graduation classification.",
        "Provide personalized academic guidance based on the authenticated student's department, batch, results, attendance, and completed credits.",
        "Calculate GPA and CGPA targets accurately using deterministic formulas rather than LLM-generated arithmetic.",
        "Explain academic risk early, including low GPA trend, missing results, low attendance, prerequisite issues, and unrealistic target GPA scenarios.",
        "Support faculty and admin users with controlled, role-based insights without exposing private student records.",
    ]
    for item in goals:
        add_bullet(doc, item)

    add_heading(doc, "4. Recommended Architecture", 1)
    add_para(
        doc,
        "The assistant should not be a single prompt connected directly to a PDF. It should be a set of controlled services where each part has a clear responsibility. The LLM should write the final explanation, but it should receive facts from trusted tools and retrieved handbook passages.",
    )
    add_table(
        doc,
        ["Layer", "Responsibility", "Implementation Notes"],
        [
            ("Chat UI", "Student and staff interface for asking academic questions.", "React/TypeScript screen inside PES with conversation history, loading states, citations, and suggested follow-up questions."),
            ("AI Gateway API", "Receives user messages, checks role, builds context, and coordinates tools.", "Server-side route or edge function; never expose service role keys to the frontend."),
            ("Intent Classifier", "Detects whether the question is about curriculum, GPA planning, attendance, handbook rules, results, or admin analytics.", "Small prompt-based classifier or rules plus embeddings; returns intent and required tools."),
            ("Retrieval Layer", "Finds handbook passages relevant to the question.", "RAG over handbook chunks with metadata such as department, semester, course code, rule type, and page number."),
            ("Supabase Tool Layer", "Fetches authenticated student facts.", "Uses RLS-safe queries/functions for profile, results, attendance, enrollments, curriculum, and GPA summaries."),
            ("Calculation Engine", "Performs GPA, CGPA, target GPA, and eligibility calculations.", "Pure backend functions with unit tests; the LLM receives the computed result only."),
            ("Answer Composer", "Produces a clear, friendly final response.", "LLM writes the answer using retrieved handbook evidence and tool outputs; must show assumptions and limitations."),
        ],
        [1350, 3000, 5010],
    )

    add_heading(doc, "5. Data Foundation", 1)
    add_para(
        doc,
        "The assistant's quality depends more on data modeling than on the model alone. Course names and codes from the handbook should be normalized into database tables. The PDF can still be used for explanations, but the application should not rely on PDF extraction every time a student asks for next semester subjects.",
    )
    add_table(
        doc,
        ["Data Area", "Suggested Tables", "Why It Matters"],
        [
            ("Student identity", "students, batches, departments", "Connects registration number, index number, email, batch, and department. For example, EN108953 maps naturally to en108953@foe.sjp.ac.lk."),
            ("Curriculum", "courses, curriculum_semesters, curriculum_courses", "Lets the assistant answer next semester subjects, credits, compulsory/elective status, and GPA contribution precisely."),
            ("Performance", "results, semester_gpa, cgpa_snapshots", "Supports result explanations, trend analysis, GPA target planning, and degree classification checks."),
            ("Attendance", "attendance_records, attendance_summaries", "Supports exam eligibility warnings and early intervention."),
            ("Handbook knowledge", "handbook_chunks, handbook_sources", "Stores retrieved passages with page, section, department, and topic metadata for cited rule explanations."),
            ("Assistant operations", "assistant_conversations, assistant_messages, assistant_feedback", "Enables auditability, feedback collection, improvement, and hallucination review."),
        ],
        [1550, 2700, 5110],
    )

    add_heading(doc, "6. Knowledge Base Strategy", 1)
    add_para(
        doc,
        "The Faculty Handbook 2026 should be processed in two ways. First, structured tables such as departments, course lists, semesters, credits, and GPA contribution should be extracted into relational tables. Second, explanatory sections such as grading rules, attendance eligibility, add/drop procedures, Dean's List rules, and award classifications should be chunked into a vector knowledge base for retrieval.",
    )
    for item in [
        "Chunk by real academic meaning, not by fixed character count only. A course table row, regulation paragraph, and award classification rule should remain coherent.",
        "Attach metadata to each chunk: handbook year, page number, section title, department, semester, course code, rule type, and effective academic year.",
        "Prefer structured SQL answers for course lists and GPA facts; use RAG when the student asks why a rule applies or asks for an explanation.",
        "Show source context in the answer, such as 'According to the Faculty Handbook 2026, ...' and keep page references internally for audit.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "7. Query Workflows", 1)
    add_heading(doc, "7.1 Next Semester Subjects", 2)
    add_para(
        doc,
        "For a question such as 'What are the subjects in next semester?', the assistant should not guess from the user's wording. It should authenticate the student, read the student's batch and department, determine the next academic semester from completed results or active enrollment, fetch the curriculum rows for that department and semester, and return the subjects grouped by compulsory, optional, and elective status.",
    )
    add_number(doc, "Read authenticated user profile from Supabase.")
    add_number(doc, "Identify department, batch, current semester, and completed semesters.")
    add_number(doc, "Fetch next semester curriculum rows with course code, course name, credits, GPA contribution, and prerequisite/elective status.")
    add_number(doc, "Return a concise answer with assumptions, such as 'assuming you have completed Semester 5'.")

    add_heading(doc, "7.2 Target GPA Planning", 2)
    add_para(
        doc,
        "For a question such as 'I want a 3.5 GPA, what GPAs should I get in next semesters?', the assistant should call a deterministic target planner. The LLM must not calculate this by itself because small arithmetic mistakes can mislead a student.",
    )
    add_code_block(
        doc,
        [
            "required_remaining_avg =",
            "  (target_final_gpa * total_gpa_credits - current_grade_points)",
            "  / remaining_gpa_credits",
        ],
    )
    add_para(
        doc,
        "The final response should include current CGPA, completed GPA credits, remaining GPA credits, the required average GPA across the remaining credits, and one or two realistic semester scenarios. If the required value is above 4.0, the assistant should clearly say that the target is mathematically impossible under the current assumptions and suggest the nearest possible outcome.",
    )

    add_heading(doc, "7.3 Handbook Rule Questions", 2)
    add_para(
        doc,
        "For handbook questions such as 'What is the requirement for Dean's List?' or 'How is degree class decided?', the assistant should retrieve relevant handbook chunks and answer in plain language. If the answer depends on personal data, it should combine the rule with the student's Supabase records.",
    )

    add_heading(doc, "7.4 Attendance and Eligibility", 2)
    add_para(
        doc,
        "Attendance questions should use the student's attendance records and faculty rules. The assistant should explain whether the student is safe, at risk, or below the expected threshold, while making the exact data date visible. This is important because attendance changes continuously.",
    )

    add_heading(doc, "8. Tool Functions", 1)
    add_table(
        doc,
        ["Tool Function", "Inputs", "Output"],
        [
            ("get_student_profile", "auth_user_id", "student id, index, registration number, email, batch, department"),
            ("get_next_semester_courses", "student_id or department + semester", "course list with code, name, credits, type, GPA contribution"),
            ("get_completed_results", "student_id", "results by semester with grades, grade points, credits, and attempts"),
            ("calculate_required_gpa", "student_id, target_cgpa", "required remaining average and possible semester scenarios"),
            ("get_attendance_risk", "student_id, course_code optional", "attendance percentage, threshold status, and risk explanation"),
            ("retrieve_handbook_context", "query + metadata filters", "relevant handbook chunks and source metadata"),
        ],
        [2350, 2600, 4410],
    )

    add_heading(doc, "9. Prompting and Response Rules", 1)
    for item in [
        "Use Supabase data and calculator outputs as facts; do not override them with model guesses.",
        "If the database is incomplete because results are still being entered, state exactly what data is missing before giving advice.",
        "Ask a clarifying question when a target depends on unknown assumptions, such as total remaining credits or repeated courses.",
        "Use a supportive tone, but keep academic decisions precise and evidence-based.",
        "For private academic data, answer only for the authenticated student unless the user has a verified staff/admin role.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10. Security and Privacy", 1)
    add_para(
        doc,
        "Because the assistant will handle student results and attendance, security must be treated as a core design requirement. The frontend should never call the LLM directly with unrestricted database access. All AI requests should pass through a backend API that validates authentication, role, and row-level permissions.",
    )
    add_table(
        doc,
        ["Risk", "Mitigation"],
        [
            ("Student sees another student's results", "Use Supabase RLS and backend role checks; tool functions must accept authenticated user context, not arbitrary student ids from the prompt."),
            ("LLM invents academic rules", "Use RAG with handbook chunks and require citations/grounding for regulation answers."),
            ("Wrong GPA calculation", "Move GPA math into tested backend functions and return computed values to the model."),
            ("Prompt injection from uploaded files", "Treat retrieved text as untrusted context; system prompt must prevent tool misuse and data exfiltration."),
            ("Incomplete result data", "Expose data completeness status and date of last update in personalized answers."),
        ],
        [2850, 6510],
    )

    add_heading(doc, "11. Evaluation Plan", 1)
    add_para(
        doc,
        "The assistant should be tested like an academic information system, not only like a chatbot. A small evaluation set should be created before launch, using real questions from students and expected answers based on the handbook and sample Supabase data.",
    )
    for item in [
        "Golden-answer tests: 50-100 common questions with expected answer patterns and required source sections.",
        "Calculator tests: unit tests for GPA, CGPA, SGPA, required target GPA, impossible targets, repeated courses, and non-GPA courses.",
        "Retrieval tests: check whether the correct handbook section appears in the top retrieved chunks.",
        "Privacy tests: verify that a student cannot request another student's results, index, email, attendance, or GPA.",
        "User feedback: add thumbs up/down and optional comments so wrong answers can be reviewed and fixed.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "12. Implementation Roadmap", 1)
    add_table(
        doc,
        ["Phase", "Scope", "Deliverable"],
        [
            ("Phase 1", "Data cleanup and curriculum modeling", "Departments, courses, curriculum semesters, and handbook chunks ready in Supabase."),
            ("Phase 2", "Basic handbook assistant", "Students can ask general faculty questions with handbook-grounded answers."),
            ("Phase 3", "Personalized student assistant", "Assistant answers next semester subjects, completed credits, CGPA summary, and missing result status."),
            ("Phase 4", "GPA target planner", "What-if planning for target CGPA, Dean's List, class targets, and semester scenarios."),
            ("Phase 5", "Risk and prediction layer", "Attendance risk, performance trend warnings, and later CART-based prediction once enough historical data exists."),
        ],
        [1200, 3300, 4860],
    )

    add_heading(doc, "13. Recommended MVP", 1)
    add_para(
        doc,
        "The best first version should focus on reliability, not maximum AI complexity. The MVP should answer handbook questions, show next semester subjects, summarize current academic standing, and calculate target GPA scenarios. Prediction models such as CART should be added only after the database has enough clean historical results to train and validate them.",
    )
    add_para(
        doc,
        "This keeps the biggest student value inside the first release: students can ask questions naturally and receive correct, personalized answers that previously required searching PDFs, checking spreadsheets, or manually calculating GPA.",
    )

    add_heading(doc, "14. Example Answers", 1)
    add_table(
        doc,
        ["Student Question", "Expected Assistant Behavior"],
        [
            ("What are my subjects next semester?", "Identify the student, determine next semester, fetch department curriculum, and list courses with credits and GPA contribution."),
            ("I want 3.5 CGPA. What should I get next?", "Use current results and remaining credits, calculate required GPA, show realistic scenarios, and state assumptions."),
            ("Can I get Dean's List this semester?", "Retrieve handbook rule, check semester results/completion status, and explain eligibility."),
            ("Am I at risk because of attendance?", "Use attendance summaries and threshold rules, show current percentage and risk level."),
            ("What does EE3205 mean?", "Fetch course details from structured curriculum and optionally explain department/code pattern from the handbook."),
        ],
        [2600, 6760],
    )

    add_heading(doc, "15. Conclusion", 1)
    add_para(
        doc,
        "The AI assistant can become the most valuable feature of the Performance Evaluation System if it is built as a grounded academic intelligence layer. The key decision is to separate conversation from truth: the LLM should communicate, but Supabase, handbook retrieval, and tested calculators should provide the facts. With this design, the assistant can answer student questions accurately, guide future academic planning, reduce manual confusion, and support faculty decision-making as the system grows.",
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
